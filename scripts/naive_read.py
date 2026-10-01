"""
Step 5A: read every sample document the naive way and save the raw text.

These are the simple readers most tutorials use. The goal is to see what they lose,
not to use them in the system.

Usage (from the project root):
    python -m scripts.naive_read
"""
from pathlib import Path

import pandas as pd

import config

OUTPUT_DIR = config.PROCESSED_DIR / "naive"


def read_pdf(path: Path) -> str:
    """Plain text of every page, as a PDF library returns it."""
    import pypdfium2 as pdfium

    pdf = pdfium.PdfDocument(str(path))
    try:
        pages = [pdf[index].get_textpage().get_text_range() for index in range(len(pdf))]
    finally:
        pdf.close()
    return "\n".join(pages)


def read_docx(path: Path) -> str:
    """Paragraph text only: the most common python-docx example."""
    from docx import Document

    return "\n".join(paragraph.text for paragraph in Document(str(path)).paragraphs)


def read_pptx(path: Path) -> str:
    """Text of every shape on every slide."""
    from pptx import Presentation

    lines = []
    for slide in Presentation(str(path)).slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                lines.append(shape.text_frame.text)
    return "\n".join(lines)


def read_html(path: Path) -> str:
    """All visible text of the page."""
    from bs4 import BeautifulSoup

    return BeautifulSoup(path.read_text(encoding="utf-8"), "html.parser").get_text("\n", strip=True)


def read_excel(path: Path) -> str:
    """pandas defaults: first sheet only, first row used as the header."""
    return pd.read_excel(path).to_string()


def read_csv(path: Path) -> str:
    """pandas defaults: comma separator, UTF-8."""
    return pd.read_csv(path).to_string()


READERS = {
    ".pdf": read_pdf,
    ".docx": read_docx,
    ".pptx": read_pptx,
    ".html": read_html,
    ".htm": read_html,
    ".xlsx": read_excel,
    ".csv": read_csv,
}


def main() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    files = sorted(path for path in config.SAMPLES_DIR.iterdir() if path.is_file())
    print(f"Reading {len(files)} files from {config.SAMPLES_DIR}\n")
    print(f"{'File':50s} {'Chars':>7s} {'Lines':>6s}  Note")

    for path in files:
        reader = READERS.get(path.suffix.lower())
        if reader is None:
            print(f"{path.name:50s} {'-':>7s} {'-':>6s}  unsupported format, skipped")
            continue
        try:
            text = reader(path)
        except Exception as exc:
            print(f"{path.name:50s} {'-':>7s} {'-':>6s}  ERROR: {type(exc).__name__}: {exc}")
            continue

        (OUTPUT_DIR / f"{path.name}.txt").write_text(text, encoding="utf-8")
        note = "EMPTY - nothing was extracted!" if not text.strip() else ""
        print(f"{path.name:50s} {len(text):>7d} {text.count(chr(10)) + 1:>6d}  {note}")

    print(f"\nOutputs written to {OUTPUT_DIR}")


if __name__ == "__main__":
    main()