"""
Step 6: edge cases for the tabular reader, using small files created on the fly.

The sample documents only contain the cases we already thought of. These tests cover variations
that real files often have: another encoding and delimiter, line breaks inside cells, number
and date formats, merged cells, columns without a header and files that cannot be opened.

Usage (from the project root):
    python -m scripts.check_tabular
"""
import sys
import tempfile
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook

from rag.tabular import TabularReadError, read_csv, read_excel
from scripts.check_env import run_check

# The test data is Turkish on purpose: these are the characters and formats real files contain.


def expect(text: str, present: list[str], absent: list[str] | None = None) -> str:
    """Pass if every 'present' phrase is in the text and no 'absent' phrase is."""
    absent = absent or []
    missing = [phrase for phrase in present if phrase not in text]
    leaked = [phrase for phrase in absent if phrase in text]
    if missing or leaked:
        raise AssertionError(f"missing {missing}, unexpected {leaked}\n--- output ---\n{text}")
    return f"{len(present) + len(absent)} expectations met"


def check_legacy_csv(folder: Path) -> str:
    """Turkish Windows encoding, comma delimiter, a quoted comma and a line break inside a cell."""
    path = folder / "legacy.csv"
    content = 'Firma,Şehir,Açıklama\n"Yıldız Ltd. Şti., Ankara",Çorum,"iki\nsatır"\n'
    path.write_bytes(content.encode("cp1254"))
    return expect(read_csv(path), [
        "Firma: Yıldız Ltd. Şti., Ankara | Şehir: Çorum",
        "Açıklama: iki satır",
    ])


def check_single_column_csv(folder: Path) -> str:
    """No delimiter to detect: must not crash, lines become paragraphs."""
    path = folder / "single.csv"
    path.write_text("Duyurular\nOfis 1 Ocak'ta kapalıdır.\n", encoding="utf-8")
    return expect(read_csv(path), ["Ofis 1 Ocak'ta kapalıdır."])


def check_excel_formats(folder: Path) -> str:
    """Number formats, dates, a merged title, a line break in a cell and a column without a header."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Bütçe"
    sheet.merge_cells("A1:F1")
    sheet["A1"] = "2026 BÜTÇE TABLOSU"
    for column, name in enumerate(["Kalem", "Yıl", "Tutar", "Pay", "Tarih"], start=1):
        sheet.cell(row=3, column=column, value=name)
    rows = [
        ["Kira", 2026, 1234.5, 0.15, datetime(2026, 7, 1), "ek bilgi"],
        ["Lisans\nyenileme", 2026, -50000, 0.085],
    ]
    for row_index, row in enumerate(rows, start=4):
        for column, value in enumerate(row, start=1):
            sheet.cell(row=row_index, column=column, value=value)
    sheet["C4"].number_format = "#,##0.00"
    sheet["D4"].number_format = "0%"
    sheet["C5"].number_format = '#,##0 "TL"'
    sheet["D5"].number_format = "0.0%"
    path = folder / "budget.xlsx"
    workbook.save(path)

    text = read_excel(path)[0]
    if text.count("2026 BÜTÇE TABLOSU") != 1:
        raise AssertionError(f"Merged title should appear exactly once:\n{text}")
    return expect(text, [
        "## Bütçe",
        "Kalem: Kira | Yıl: 2026 | Tutar: 1.234,50 | Pay: %15 | Tarih: 01.07.2026 | Column 6: ek bilgi",
        "Kalem: Lisans yenileme | Yıl: 2026 | Tutar: -50.000 TL | Pay: %8,5",
    ], absent=["2.026", "00:00"])


def check_broken_excel(folder: Path) -> str:
    """A file that is not a real workbook (or is password-protected) must fail loudly."""
    path = folder / "broken.xlsx"
    path.write_bytes(b"this is not an Excel file")
    try:
        read_excel(path)
    except TabularReadError as exc:
        return f"clear error: {exc}"
    raise AssertionError("A broken file was read without an error.")


def main() -> None:
    checks = {
        "CSV in Turkish Windows encoding": check_legacy_csv,
        "CSV with a single column": check_single_column_csv,
        "Excel number and date formats": check_excel_formats,
        "Broken Excel file": check_broken_excel,
    }
    with tempfile.TemporaryDirectory() as tmp:
        folder = Path(tmp)
        results = [run_check(name, lambda check=check: check(folder)) for name, check in checks.items()]
    print(f"\n{sum(results)}/{len(results)} checks passed.")
    sys.exit(0 if all(results) else 1)


if __name__ == "__main__":
    main()